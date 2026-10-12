import io
import pytest
from docx import Document
from reportlab.pdfgen import canvas
from studio.ingest import extract_text

SRT = "1\n00:00:01,000 --> 00:00:03,000\nXin chào\n\n2\n00:00:03,500 --> 00:00:05,000\nThế giới\n"

def test_txt():
    assert extract_text("a.txt", "Lạm phát là gì?".encode()) == "Lạm phát là gì?"

def test_cp1258_not_decoded_as_utf16():
    text = "Xin chào ban đ"
    data = text.encode("cp1258")
    assert len(data) % 2 == 0  # even length: the old utf-16 path would accept it
    assert extract_text("a.txt", data) == text

def test_utf16_with_bom():
    assert extract_text("a.txt", "Xin chào".encode("utf-16")) == "Xin chào"

def test_srt_strips_timestamps():
    assert extract_text("a.srt", SRT.encode()) == "Xin chào\nThế giới"

def test_docx():
    d = Document(); d.add_paragraph("Đoạn một"); d.add_paragraph("Đoạn hai")
    buf = io.BytesIO(); d.save(buf)
    assert extract_text("a.docx", buf.getvalue()) == "Đoạn một\nĐoạn hai"

def test_pdf():
    buf = io.BytesIO(); c = canvas.Canvas(buf); c.drawString(72, 720, "Hello PDF"); c.save()
    assert "Hello PDF" in extract_text("a.pdf", buf.getvalue())

def test_unsupported():
    with pytest.raises(ValueError):
        extract_text("a.exe", b"x")
